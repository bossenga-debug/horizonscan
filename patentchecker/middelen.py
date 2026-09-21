# -*- coding: utf-8 -*-
"""Welke middelen de patentchecker volgt, en onder welke naam.

Gedeeld door ophalen.py (die per middel het octrooiregister en ClinicalTrials.gov
bevraagt) en bouw_site.py (die alles samenvoegt). Zo kan de selectie nooit
tussen die twee uit elkaar lopen.

De selectie: elk add-on geneesmiddel (ATC-5) waarvan de Zvw-vergoeding in een
van de twee laatste GIP-jaren boven DREMPEL lag. Twee jaren en niet één, omdat
het laatste jaar voorlopig is en nog niet volledig gedeclareerd: een middel dat
in 2024 € 11 mln kostte en in het halfvolle 2025 € 9 mln, hoort er gewoon bij.
"""
import os
import re
import unicodedata
import zipfile

HIER = os.path.dirname(os.path.abspath(__file__))
BRON = os.path.join(HIER, 'bron')

DREMPEL = 10_000_000

# ATC-codes die geen enkelvoudige werkzame stof zijn maar een productgroep. Een
# SPC hoort bij één stof, dus hier valt geen patentdatum te geven. Ze blijven
# wél in het overzicht staan: ze horen bij de duurste middelen, en een lege
# plek zonder uitleg zou lijken op een fout.
GROEPEN = {
    'J06BA02': 'Plasmaproduct: meerdere merken van verschillende donorpools; '
               'geen SPC en geen biosimilarroute.',
    'B02BD02': 'Stollingsfactoren VIII: een groep van recombinante en plasmaproducten, '
               'elk met een eigen INN. Ook de Horizonscan laat stollingsfactoren buiten '
               'het patentoverzicht.',
    'B02BD04': 'Stollingsfactoren IX: een groep van recombinante en plasmaproducten, '
               'elk met een eigen INN. Ook de Horizonscan laat stollingsfactoren buiten '
               'het patentoverzicht.',
    'M03AX01': 'Botulinetoxine: meerdere producten (type A en B) die niet onderling '
               'uitwisselbaar zijn; de oorspronkelijke octrooien zijn lang verlopen.',
}

# De Engelse INN waaronder het middel bij EMA en in het octrooiregister staat,
# waar die niet vanzelf uit de GIP-naam volgt (zie inn_voor()).
INN = {
    'L01XL03': 'axicabtagene ciloleucel',
    'H01AC01': 'somatropin',
    'L02BX03': 'abiraterone',
}

# Extra zoektermen voor het octrooiregister. Het register zoekt in de titel van
# het certificaat, en die is soms een merknaam of een chemische naam in plaats
# van de INN. Wat hier niet staat en toch niet gevonden wordt, valt terug op de
# datum uit de Horizonscan-PDF -- dat is op de pagina te zien.
RVO_EXTRA = {
    'L01EB04': ['Tagrisso'],
    'L01EC02': ['Tafinlar'],
    'L01XL03': ['Yescarta', 'axicabtagen'],
    'L04AB02': ['Remicade'],
    'M09AX09': ['Zolgensma'],
}


def sleutel(tekst):
    """Vergelijkingssleutel die Nederlandse en Engelse stofnamen gelijk maakt.

    GIP schrijft 'Darbepoetine alfa' en 'Onasemnogeen', EMA 'darbepoetin alfa'
    en 'onasemnogene'. Per woord de slot-e weg en ee -> e maakt die gelijk,
    zonder te gaan gokken: twee verschillende stoffen krijgen zo nooit dezelfde
    sleutel (anders dan bij fuzzy matchen, waar deuruxolitinib bij ruxolitinib
    uitkomt).
    """
    plat = unicodedata.normalize('NFKD', tekst or '').encode('ascii', 'ignore').decode().lower()
    woorden = re.findall(r'[a-z0-9]+', plat)
    return ' '.join(re.sub(r'e$', '', w.replace('ee', 'e')) for w in woorden)


# Zoutvormen die in de INN-kolom van EMA achter de stofnaam kunnen staan.
# 'trastuzumab emtansine' is géén zoutvorm van trastuzumab -- daarom een lijst en
# geen 'alles wat met de naam begint'.
ZOUTEN = {'hydrochloride', 'dihydrochloride', 'mesylate', 'mesilate', 'phosphate',
          'hemifumarate', 'fumarate', 'maleate', 'acetate', 'citrate', 'tartrate',
          'besilate', 'sulfate', 'sodium', 'potassium', 'calcium', 'tosylate',
          'succinate', 'hydrobromide', 'monohydrate', 'isethionate', 'etexilate'}


def zelfde_stof(ema_inn, inn):
    a, b = (ema_inn or '').lower().strip(), inn.lower().strip()
    if a == b:
        return True
    if a.startswith(b + ' '):
        rest = a[len(b) + 1:].split()
        return bool(rest) and all(w in ZOUTEN for w in rest)
    return False


def lees_gip(pad=None):
    """GIP add-on Zvw meerjaren -> {atc: {'naam', 'kosten': {jaar: €}, 'gebruikers': {...}}}

    Het bestand gebruikt '#' als scheidingsteken en zet een '*' achter voorlopige
    jaren; die ster bewaren we apart."""
    pad = pad or os.path.join(BRON, 'gip_addon.csv')
    uit, voorlopig = {}, set()
    with open(pad, encoding='utf-8-sig') as f:
        kop = [k.strip() for k in f.readline().split('#')]
        if kop[:4] != ['jaar', 'atclaatst', 'atclaatst_naam_tekst', 'vergoeding']:
            raise RuntimeError(f'GIP-bestand heeft een onverwachte kop: {kop}')
        for regel in f:
            v = [x.strip() for x in regel.split('#')]
            if len(v) < 5 or not v[1]:
                continue
            jaar = v[0].rstrip('*')
            if v[0].endswith('*'):
                voorlopig.add(jaar)
            o = uit.setdefault(v[1], {'naam': ' '.join(v[2].split()), 'kosten': {}, 'gebruikers': {}})
            o['kosten'][jaar] = float(v[3] or 0)
            o['gebruikers'][jaar] = int(float(v[4] or 0))
    return uit, sorted(voorlopig)


def selectie(gip):
    jaren = sorted({j for o in gip.values() for j in o['kosten']})
    laatste = jaren[-2:]
    gekozen = [a for a, o in gip.items() if max(o['kosten'].get(j, 0) for j in laatste) >= DREMPEL]
    return sorted(gekozen, key=lambda a: -max(gip[a]['kosten'].get(j, 0) for j in laatste)), jaren


def lees_ema(pad=None):
    """EMA-medicijnentabel -> lijst dicts (alleen humane middelen).

    De kop staat niet op de eerste regel (er gaan een paar regels met
    exportinformatie aan vooraf), dus we zoeken hem op."""
    import openpyxl
    pad = pad or os.path.join(BRON, 'ema_medicines.xlsx')
    rijen = openpyxl.load_workbook(pad, read_only=True).active.iter_rows(values_only=True)
    kop = None
    uit = []
    for r in rijen:
        if kop is None:
            if r and r[0] == 'Category':
                kop = [str(k or '').replace('\n', ' ').strip() for k in r]
                for nodig in ('Name of medicine', 'Medicine status', 'Biosimilar', 'Generic',
                              'International non-proprietary name (INN) / common name',
                              'Marketing authorisation date', 'Orphan medicine'):
                    if nodig not in kop:
                        raise RuntimeError(f'EMA-tabel: kolom "{nodig}" ontbreekt')
            continue
        if r and r[0] == 'Human':
            uit.append(dict(zip(kop, r)))
    if kop is None:
        raise RuntimeError('EMA-tabel: kopregel niet gevonden')
    return uit


def inn_voor(atc, gip_naam, ema_inns):
    """De Engelse INN bij een GIP-regel: handmatig, anders via de sleutel."""
    if atc in INN:
        return INN[atc]
    k = sleutel(gip_naam)
    for inn in ema_inns:
        if sleutel(inn) == k:
            return inn
    return gip_naam.lower()


def lees_gs(pad=None):
    """Farmatec add-on GS-lijst -> {atc: {'reg': {EU-nummer: fabrikant}, 'artikelen': [...]}}

    Zelfde lezing als het add-on dashboard: voor concurrentie tellen de
    handelsvergunningen, niet de kolom fabrikant (die telt parallelimport mee)."""
    import openpyxl
    pad = pad or os.path.join(BRON, 'addon-gs.zip')
    zf = zipfile.ZipFile(pad)
    naam = next(n for n in zf.namelist() if n.lower().endswith('.xlsx'))
    maand = re.search(r'(januari|februari|maart|april|mei|juni|juli|augustus|september|'
                      r'oktober|november|december)[-_ ]?(\d{4})', naam, re.I)
    it = openpyxl.load_workbook(zf.open(naam), read_only=True, data_only=True).active.iter_rows(values_only=True)
    next(it)
    uit = {}
    for r in it:
        if r[0] is None:
            continue
        zi, art, code, stof, regnr, fab = r[:6]
        code = (code or '').replace('_', '').strip()
        if not code:
            continue
        o = uit.setdefault(code, {'reg': {}, 'artikelen': set()})
        m = re.match(r'(EU/\d+/\d+/\d+)', str(regnr or '').strip())
        reg = m.group(1) if m else str(regnr or '').strip()[:14]
        if reg:
            o['reg'].setdefault(reg, str(fab or '').strip())
        merk = re.split(r'\s+\d', str(art or '').strip())[0].strip().title()
        if merk:
            o['artikelen'].add(merk)
    versie = maand.group(0).replace('-', ' ').replace('_', ' ') if maand else ''
    return {a: {'reg': o['reg'], 'artikelen': sorted(o['artikelen'])} for a, o in uit.items()}, versie
